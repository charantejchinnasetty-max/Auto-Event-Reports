"""
Daily attendance report for a class event (Dr. Bhatia Medical Coaching Institute).

Input : the Excel exported from Class Events > View Students > Export
        (contains ALL scans for the event so far).
Output: a formatted Excel with
        - Daily       : students scanned per batch per day + total
        - Cumulative  : running total per batch up to each day
        Plus unique-student and total-scan counts.

Usage : python event_report.py export.xlsx report.xlsx "SURGERY PART-2(B)TS"   (Excel)
        python event_report.py export.xlsx report.png  "SURGERY PART-2(B)TS"   (image)
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# If the panel has inconsistent batch names for the SAME batch, map them here.
BATCH_ALIASES = {
    # "SSP2 (B) 2026 ACHIEVERS BATCH": "SSP2 (B) T.S 2026 ACHIEVERS BATCH",
}
COUNT_ONLY_STATUS = None   # e.g. "Active" to count only that status; None = all rows


def _parse_dates(col):
    """Handles '3-10-2026' (day first), '2026-10-03', '2026-10-03 00:00:00' and real Excel dates."""
    s = col.astype(str).str.strip()
    iso = s.str.match(r"^\d{4}-\d{2}-\d{2}")
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    out[iso] = pd.to_datetime(s[iso].str[:10], format="%Y-%m-%d", errors="coerce")
    out[~iso] = pd.to_datetime(s[~iso], dayfirst=True, errors="coerce")
    return out.dt.normalize()


def load(path):
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)       # every sheet
    df = None
    for name, d in sheets.items():                                  # use the sheet that has the data
        d.columns = [str(c).strip() for c in d.columns]
        if {"Batch", "Scanned_date"} <= set(d.columns):
            df = d
            break
    if df is None:
        found = list(next(iter(sheets.values())).columns) if sheets else []
        raise KeyError(f"Batch / Scanned_date (columns found: {found})")
    df = df.copy()
    df["Batch"] = (df["Batch"].fillna("(no batch)").str.split().str.join(" ")
                   .replace("", "(no batch)").replace(BATCH_ALIASES))   # tabs / double spaces
    df["Date"] = _parse_dates(df["Scanned_date"])
    bad = int(df["Date"].isna().sum())
    df = df[df["Date"].notna()]
    if df.empty:
        raise ValueError("No valid Scanned_date values found in this file.")
    df.attrs["skipped_rows"] = bad
    if "Admission_no" not in df.columns:
        df["Admission_no"] = df.get("Student_dynamic_id", pd.Series(range(len(df)), index=df.index))
    if COUNT_ONLY_STATUS:
        df = df[df["Status"].str.strip() == COUNT_ONLY_STATUS]
    return df


def build(df, out_path, event_name):
    days = pd.date_range(df["Date"].min(), df["Date"].max())
    daily = (df.groupby(["Batch", "Date"]).size().unstack(fill_value=0)
               .reindex(columns=days, fill_value=0))
    # biggest batch first
    daily = daily.loc[daily.sum(axis=1).sort_values(ascending=False).index]
    uniq_per_day = df.groupby("Date")["Admission_no"].nunique().reindex(days, fill_value=0)

    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3864")
    tot_fill = PatternFill("solid", fgColor="D9E1F2")
    f = lambda **k: Font(name="Arial", size=10, **k)

    wb = Workbook()

    def sheet(ws, title, cumulative):
        ws["A1"] = f"{event_name} - {title}"
        ws["A1"].font = Font(name="Arial", size=13, bold=True)
        ws["A2"] = f"Report date: {days[-1]:%d-%m-%Y}   |   Event day: {len(days)}"
        ws["A2"].font = f(italic=True)

        hdr = ["Batch"] + [f"Day {i+1}\n{d:%d-%m-%Y}" for i, d in enumerate(days)]
        hdr += ["Total"] if not cumulative else []
        r0 = 4
        for c, h in enumerate(hdr, 1):
            cell = ws.cell(r0, c, h)
            cell.font = f(bold=True, color="FFFFFF"); cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = box
        ws.row_dimensions[r0].height = 32

        n = len(days)
        for i, (batch, row) in enumerate(daily.iterrows()):
            r = r0 + 1 + i
            ws.cell(r, 1, batch)
            for j, d in enumerate(days):
                col = 2 + j
                if cumulative:   # running total of the Daily sheet
                    L = get_column_letter(col)
                    ws.cell(r, col, f"=SUM(Daily!$B{r}:{L}{r})")
                else:
                    ws.cell(r, col, int(row[d]))
            if not cumulative:
                a, b = get_column_letter(2), get_column_letter(1 + n)
                ws.cell(r, 2 + n, f"=SUM({a}{r}:{b}{r})")
        last = r0 + len(daily)
        tr = last + 1
        ws.cell(tr, 1, "Grand Total")
        for col in range(2, 2 + n + (0 if cumulative else 1)):
            L = get_column_letter(col)
            ws.cell(tr, col, f"=SUM({L}{r0+1}:{L}{last})")
        for r in range(r0 + 1, tr + 1):
            for c in range(1, len(hdr) + 1):
                cell = ws.cell(r, c)
                cell.font = f(bold=(r == tr)); cell.border = box
                if c > 1: cell.alignment = Alignment(horizontal="center")
                if r == tr: cell.fill = tot_fill
        ws.column_dimensions["A"].width = 40
        for c in range(2, len(hdr) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.freeze_panes = ws.cell(r0 + 1, 2)
        return tr

    ws1 = wb.active; ws1.title = "Daily"
    tr = sheet(ws1, "Daily Attendance (scans per day)", cumulative=False)

    # Unique students block under the Daily table
    u = tr + 2
    ws1.cell(u, 1, "Unique students (distinct admission no.)").font = f(bold=True)
    for j, d in enumerate(days):
        ws1.cell(u, 2 + j, int(uniq_per_day[d])).font = f()
        ws1.cell(u, 2 + j).alignment = Alignment(horizontal="center")
    ws1.cell(u, 2 + len(days), int(df["Admission_no"].nunique())).font = f(bold=True)
    ws1.cell(u, 2 + len(days)).alignment = Alignment(horizontal="center")
    ws1.cell(u + 1, 1,
        "Note: 'Grand Total' counts every scan (a student present on 2 days counts twice). "
        "'Unique students' counts each student once.").font = f(italic=True, color="666666")

    ws2 = wb.create_sheet("Cumulative")
    sheet(ws2, "Cumulative Attendance (running total)", cumulative=True)

    wb.save(out_path)


def build_image(df, out_path, event_name):
    """Render the Daily + Cumulative tables as one PNG (easy to send on WhatsApp/email)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    days = pd.date_range(df["Date"].min(), df["Date"].max())
    daily = (df.groupby(["Batch", "Date"]).size().unstack(fill_value=0)
               .reindex(columns=days, fill_value=0))
    daily = daily.loc[daily.sum(axis=1).sort_values(ascending=False).index]
    cum = daily.cumsum(axis=1)
    uniq_day = df.groupby("Date")["Admission_no"].nunique().reindex(days, fill_value=0)

    heads = [f"Day {i+1}\n{d:%d-%m}" for i, d in enumerate(days)]

    def table(ax, title, data, total_col, extra_row=None):
        ax.axis("off")
        ax.set_title(title, loc="left", fontsize=12, fontweight="bold", pad=8)
        cols = ["Batch"] + heads + (["Total"] if total_col else [])
        rows = []
        for b, r in data.iterrows():
            rows.append([b] + [int(v) for v in r] + ([int(r.sum())] if total_col else []))
        tot = ["Grand Total"] + [int(data[d].sum()) for d in days] + ([int(data.values.sum())] if total_col else [])
        rows.append(tot)
        if extra_row:
            rows.append(extra_row)
        t = ax.table(cellText=rows, colLabels=cols, loc="upper left", cellLoc="center")
        t.auto_set_font_size(False); t.set_fontsize(10); t.scale(1, 1.7)
        t.auto_set_column_width(list(range(len(cols))))
        for (r, c), cell in t.get_celld().items():
            cell.set_edgecolor("#BBBBBB")
            if r == 0:
                cell.set_facecolor("#1F3864"); cell.get_text().set_color("white")
                cell.get_text().set_fontweight("bold"); cell.set_height(cell.get_height() * 1.4)
            elif r == len(data) + 1:
                cell.set_facecolor("#D9E1F2"); cell.get_text().set_fontweight("bold")
            elif extra_row and r == len(data) + 2:
                cell.set_facecolor("#F2F2F2")
            if c == 0 and r > 0:
                cell.get_text().set_ha("left"); cell._loc = "left"

    n1, n2 = len(daily) + 3, len(daily) + 2          # table rows incl. header/total/extra
    fig, axes = plt.subplots(2, 1, figsize=(7.5 + 0.9 * len(days), 0.42 * (n1 + n2) + 2.0),
                             gridspec_kw={"height_ratios": [n1, n2], "hspace": 0.25})
    fig.suptitle(f"{event_name}\nReport date: {days[-1]:%d-%m-%Y}  |  Event day {len(days)}",
                 x=0.02, ha="left", fontsize=14, fontweight="bold")
    unique_row = ["Unique students"] + [int(v) for v in uniq_day] + [int(df["Admission_no"].nunique())]
    table(axes[0], "Daily attendance (scans per day)", daily, True, unique_row)
    table(axes[1], "Cumulative attendance (running total)", cum, False)
    fig.subplots_adjust(left=0.02, right=0.98, top=0.84, bottom=0.02)
    fig.savefig(out_path, dpi=160, facecolor="white")
    plt.close(fig)


REQUIRED = ["Batch", "Scanned_date"]


def summarize(df):
    """JSON-friendly summary used by the web page."""
    days = pd.date_range(df["Date"].min(), df["Date"].max())
    daily = (df.groupby(["Batch", "Date"]).size().unstack(fill_value=0)
               .reindex(columns=days, fill_value=0))
    daily = daily.loc[daily.sum(axis=1).sort_values(ascending=False).index]
    cum = daily.cumsum(axis=1)
    uniq = df.groupby("Date")["Admission_no"].nunique().reindex(days, fill_value=0)
    return {
        "days": [d.strftime("%d-%m-%Y") for d in days],
        "batches": [{"batch": b, "daily": [int(v) for v in daily.loc[b]],
                     "cumulative": [int(v) for v in cum.loc[b]],
                     "total": int(daily.loc[b].sum())} for b in daily.index],
        "total_daily": [int(v) for v in daily.sum()],
        "total_cumulative": [int(v) for v in cum.sum()],
        "grand_total": int(daily.values.sum()),
        "unique_daily": [int(v) for v in uniq],
        "unique_overall": int(df["Admission_no"].nunique()),
    }


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    name = sys.argv[3] if len(sys.argv) > 3 else "Class Event"
    data = load(src)
    if dst.lower().endswith(".png"):
        build_image(data, dst, name)
    else:
        build(data, dst, name)
    print("Saved", dst)
