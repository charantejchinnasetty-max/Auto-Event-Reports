"""
Daily attendance report for a class event (Dr. Bhatia Medical Coaching Institute).

Input : the Excel exported from Class Events > View Students > Export
        (contains ALL scans for the event so far).
Output: a formatted Excel with
        - Daily       : students scanned per batch per day + total
        - Cumulative  : running total per batch up to each day
        Plus unique-student and total-scan counts.

Usage : python event_report.py export.xlsx report.xlsx "SURGERY PART-2(B)TS"   (Excel)
        python event_report.py export.xlsx report.png  "SURGERY PART-2(B)TS"   (image)
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# If the panel has inconsistent batch names for the SAME batch, map them here.
BATCH_ALIASES = {
    # "SSP2 (B) 2026 ACHIEVERS BATCH": "SSP2 (B) T.S 2026 ACHIEVERS BATCH",
}
COUNT_ONLY_STATUS = None   # e.g. "Active" to count only that status; None = all rows


def _parse_dates(col):
    """Handles '3-10-2026' (day first), '2026-10-03', '2026-10-03 00:00:00' and real Excel dates."""
    s = col.astype(str).str.strip()
    iso = s.str.match(r"^\d{4}-\d{2}-\d{2}")
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    out[iso] = pd.to_datetime(s[iso].str[:10], format="%Y-%m-%d", errors="coerce")
    out[~iso] = pd.to_datetime(s[~iso], format="mixed", dayfirst=True, errors="coerce")
    return out.dt.normalize()


def load(path):
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)       # every sheet
    df = None
    for name, d in sheets.items():                                  # use the sheet that has the data
        d.columns = [str(c).strip() for c in d.columns]
        if {"Batch", "Scanned_date"} <= set(d.columns):
            df = d
            break
    if df is None:
        found = list(next(iter(sheets.values())).columns) if sheets else []
        raise KeyError(f"Batch / Scanned_date (columns found: {found})")
    df = df.copy()
    df["Batch"] = (df["Batch"].fillna("(no batch)").str.split().str.join(" ")
                   .replace("", "(no batch)").replace(BATCH_ALIASES))   # tabs / double spaces
    if "Admission_no" in df.columns:
        df["Admission_no"] = df["Admission_no"].str.strip().replace("", pd.NA)
    df["Date"] = _parse_dates(df["Scanned_date"])
    bad = int(df["Date"].isna().sum())
    df = df[df["Date"].notna()]
    if df.empty:
        raise ValueError("No valid Scanned_date values found in this file.")
    df.attrs["skipped_rows"] = bad
    if "Admission_no" not in df.columns:
        if "Student_dynamic_id" not in df.columns:
            raise ValueError("Missing student identifier: Admission_no or Student_dynamic_id is required.")
        df["Admission_no"] = df["Student_dynamic_id"].str.strip().replace("", pd.NA)
    if COUNT_ONLY_STATUS:
        df = df[df["Status"].str.strip() == COUNT_ONLY_STATUS]
    return df


def build(df, out_path, event_name):
    days = pd.date_range(df["Date"].min(), df["Date"].max())
    daily = (df.groupby(["Batch", "Date"]).size().unstack(fill_value=0)
               .reindex(columns=days, fill_value=0))
    # biggest batch first
    daily = daily.loc[daily.sum(axis=1).sort_values(ascending=False).index]
    uniq_per_day = df.groupby("Date")["Admission_no"].nunique().reindex(days, fill_value=0)

    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3864")
    tot_fill = PatternFill("solid", fgColor="D9E1F2")
    f = lambda **k: Font(name="Arial", size=10, **k)

    wb = Workbook()

    def sheet(ws, title, cumulative):
        ws["A1"] = f"{event_name} - {title}"
        ws["A1"].font = Font(name="Arial", size=13, bold=True)
        ws["A2"] = f"Report date: {days[-1]:%d-%m-%Y}   |   Event day: {len(days)}"
        ws["A2"].font = f(italic=True)

        hdr = ["Batch"] + [f"Day {i+1}\n{d:%d-%m-%Y}" for i, d in enumerate(days)]
        hdr += ["Total"] if not cumulative else []
        r0 = 4
        for c, h in enumerate(hdr, 1):
            cell = ws.cell(r0, c, h)
            cell.font = f(bold=True, color="FFFFFF"); cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = box
        ws.row_dimensions[r0].height = 32

        n = len(days)
        for i, (batch, row) in enumerate(daily.iterrows()):
            r = r0 + 1 + i
            ws.cell(r, 1, batch)
            for j, d in enumerate(days):
                col = 2 + j
                if cumulative:   # running total of the Daily sheet
                    L = get_column_letter(col)
                    ws.cell(r, col, f"=SUM(Daily!$B{r}:{L}{r})")
                else:
                    ws.cell(r, col, int(row[d]))
            if not cumulative:
                a, b = get_column_letter(2), get_column_letter(1 + n)
                ws.cell(r, 2 + n, f"=SUM({a}{r}:{b}{r})")
        last = r0 + len(daily)
        tr = last + 1
        ws.cell(tr, 1, "Grand Total")
        for col in range(2, 2 + n + (0 if cumulative else 1)):
            L = get_column_letter(col)
            ws.cell(tr, col, f"=SUM({L}{r0+1}:{L}{last})")
        for r in range(r0 + 1, tr + 1):
            for c in range(1, len(hdr) + 1):
                cell = ws.cell(r, c)
                cell.font = f(bold=(r == tr)); cell.border = box
                if c > 1: cell.alignment = Alignment(horizontal="center")
                if r == tr: cell.fill = tot_fill
        ws.column_dimensions["A"].width = 40
        for c in range(2, len(hdr) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.freeze_panes = ws.cell(r0 + 1, 2)
        return tr

    ws1 = wb.active; ws1.title = "Daily"
    tr = sheet(ws1, "Daily Attendance (scans per day)", cumulative=False)

    # Unique students block under the Daily table
    u = tr + 2
    ws1.cell(u, 1, "Unique students (distinct admission no.)").font = f(bold=True)
    for j, d in enumerate(days):
        ws1.cell(u, 2 + j, int(uniq_per_day[d])).font = f()
        ws1.cell(u, 2 + j).alignment = Alignment(horizontal="center")
    ws1.cell(u, 2 + len(days), int(df["Admission_no"].nunique())).font = f(bold=True)
    ws1.cell(u, 2 + len(days)).alignment = Alignment(horizontal="center")
    ws1.cell(u + 1, 1,
        "Note: 'Grand Total' counts every scan (a student present on 2 days counts twice). "
        "'Unique students' counts each student once.").font = f(italic=True, color="666666")

    ws2 = wb.create_sheet("Cumulative")
    sheet(ws2, "Cumulative Attendance (running total)", cumulative=True)

    wb.save(out_path)


def build_image(df, out_path, event_name):
    """Render an adaptive, high-resolution report without browser dependencies."""
    from PIL import Image, ImageDraw, ImageFont
    import textwrap
    report = summarize(df)
    n = len(report["days"])
    width = max(1500, 600 + (n + 1) * 125)
    margin = 64
    batch_width = 500
    cell_width = (width - margin * 2 - batch_width) / (n + 1)
    font_dir = __import__('pathlib').Path('C:/Windows/Fonts')
    def font(size, bold=False):
        for name in (["arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold else ["arial.ttf", "DejaVuSans.ttf"]):
            try:
                return ImageFont.truetype(str(font_dir / name) if (font_dir / name).exists() else name, size)
            except OSError:
                pass
        return ImageFont.load_default()
    title_lines = textwrap.wrap(event_name, max(20, int((width - 128) / 29))) or ["Class Event"]
    heights = [max(58, 28 * len(textwrap.wrap(b["batch"], 35)) + 20) for b in report["batches"]]
    header_height = 190 + (len(title_lines) - 1) * 56
    height = header_height + 360 + 2 * (126 + sum(heights)) + 600
    im = Image.new("RGB", (width, height), "#eef3f8")
    draw = ImageDraw.Draw(im)
    navy, teal, muted = "#132d46", "#087e8b", "#607489"
    draw.rectangle((0, 0, width, header_height), fill=navy)
    draw.text((margin, 30), "ATTENDANCE REPORT", font=font(19, True), fill="#77d5dc")
    y = 69
    for line in title_lines:
        draw.text((margin, y), line, font=font(44, True), fill="white")
        y += 56
    draw.text((margin, y + 9), f"{report['days'][0]} to {report['days'][-1]}    /    Event day {n}", font=font(22), fill="#c7d8e8")
    y = header_height + 28
    cards = [("TOTAL SCANS", report["grand_total"]), ("UNIQUE STUDENTS", report["unique_overall"]), ("BATCHES", len(report["batches"])), ("LATEST DAY SCANS", report["total_daily"][-1])]
    cw = (width - 2 * margin - 3 * 18) / 4
    for i, (label, value) in enumerate(cards):
        x = margin + i * (cw + 18)
        draw.rounded_rectangle((x, y, x + cw, y + 113), radius=16, fill="white")
        draw.text((x + 22, y + 18), label, font=font(16, True), fill=muted)
        draw.text((x + 22, y + 47), f"{value:,}", font=font(39, True), fill=navy)
    y += 141
    draw.text((margin, y), "Daily attendance trend", font=font(24, True), fill=navy)
    y += 48
    bar_space = (width - 2 * margin) / n
    peak = max(report["total_daily"])
    for i, value in enumerate(report["total_daily"]):
        x = margin + i * bar_space
        bw = max(15, min(120, bar_space - 25))
        h = 78 * value / peak
        draw.rounded_rectangle((x, y + 85 - h, x + bw, y + 85), radius=5, fill=teal)
        draw.text((x, y + 85 - h - 26), str(value), font=font(19, True), fill=navy)
        draw.text((x, y + 93), report["days"][i][:5], font=font(16), fill=muted)
    y += 151
    def table(y, title, cumulative=False):
        draw.text((margin, y), title, font=font(27, True), fill=navy)
        y += 49
        headers = ["Batch"] + [f"Day {i+1}\n{d[:5]}" for i, d in enumerate(report["days"])] + ["Total"]
        def row(y, values, rh, fill, color=navy, bold=False):
            draw.rectangle((margin, y, width-margin, y+rh), fill=fill)
            for j, value in enumerate(values):
                left = margin if j == 0 else margin + batch_width + (j-1)*cell_width
                w = batch_width if j == 0 else cell_width
                text = "\n".join(textwrap.wrap(str(value), 35)) if j == 0 else str(value)
                f = font(19 if j == 0 else 21, bold)
                box = draw.multiline_textbbox((0, 0), text, font=f, spacing=5)
                tx = left + 18 if j == 0 else left + (w - (box[2]-box[0]))/2
                draw.multiline_text((tx, y + (rh - (box[3]-box[1]))/2 - box[1]), text, font=f, fill=color, spacing=5, align="center" if j else "left")
            draw.line((margin, y+rh, width-margin, y+rh), fill="#dfe7ef", width=1)
            return y + rh
        y = row(y, headers, 76, navy, "white", True)
        for i, b in enumerate(report["batches"]):
            vals = b["cumulative"] if cumulative else b["daily"]
            y = row(y, [b["batch"]] + vals + [b["total"]], heights[i], "white" if i%2 == 0 else "#f5f8fb")
        vals = report["total_cumulative"] if cumulative else report["total_daily"]
        y = row(y, ["Grand Total"] + vals + [report["grand_total"]], 60, "#dceef0", navy, True)
        if not cumulative:
            y = row(y, ["Unique students"] + report["unique_daily"] + [report["unique_overall"]], 58, "#e5ebf3")
        return y
    y = table(y, "01   Daily attendance", False) + 40
    y = table(y, "02   Cumulative attendance", True) + 28
    draw.text((margin, y), "Scans count every attendance record. Unique students count each admission number once.", font=font(18), fill=muted)
    if report["skipped_rows"] or report["missing_ids"]:
        y += 27
        draw.text((margin, y), f"Invalid-date rows skipped: {report['skipped_rows']}   /   Scans without a student ID: {report['missing_ids']}", font=font(18), fill=muted)
    im = im.crop((0, 0, width, y + 62))
    im.save(out_path, format="PNG")


REQUIRED = ["Batch", "Scanned_date"]


def summarize(df):
    """JSON-friendly summary used by the web page."""
    days = pd.date_range(df["Date"].min(), df["Date"].max())
    daily = (df.groupby(["Batch", "Date"]).size().unstack(fill_value=0)
               .reindex(columns=days, fill_value=0))
    daily = daily.loc[daily.sum(axis=1).sort_values(ascending=False).index]
    cum = daily.cumsum(axis=1)
    uniq = df.groupby("Date")["Admission_no"].nunique().reindex(days, fill_value=0)
    return {
        "days": [d.strftime("%d-%m-%Y") for d in days],
        "batches": [{"batch": b, "daily": [int(v) for v in daily.loc[b]],
                     "cumulative": [int(v) for v in cum.loc[b]],
                     "total": int(daily.loc[b].sum())} for b in daily.index],
        "total_daily": [int(v) for v in daily.sum()],
        "total_cumulative": [int(v) for v in cum.sum()],
        "grand_total": int(daily.values.sum()),
        "unique_daily": [int(v) for v in uniq],
        "unique_overall": int(df["Admission_no"].nunique()),
        "skipped_rows": int(df.attrs.get("skipped_rows", 0)),
        "missing_ids": int(df["Admission_no"].isna().sum()),
    }


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    name = sys.argv[3] if len(sys.argv) > 3 else "Class Event"
    data = load(src)
    if dst.lower().endswith(".png"):
        build_image(data, dst, name)
    else:
        build(data, dst, name)
    print("Saved", dst)
